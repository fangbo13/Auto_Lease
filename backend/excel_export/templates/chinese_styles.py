from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers

# 边框样式
THIN_BORDER = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000'),
)

MEDIUM_BORDER = Border(
    left=Side(style='medium', color='000000'),
    right=Side(style='medium', color='000000'),
    top=Side(style='medium', color='000000'),
    bottom=Side(style='medium', color='000000'),
)

# 字体样式
TITLE_FONT = Font(name='微软雅黑', size=14, bold=True, color='FFFFFF')
HEADER_FONT = Font(name='微软雅黑', size=11, bold=True, color='000000')
DATA_FONT = Font(name='微软雅黑', size=10, color='000000')
LABEL_FONT = Font(name='微软雅黑', size=10, bold=True, color='000000')

# 填充样式
TITLE_FILL = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
HEADER_FILL = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')
SUBHEADER_FILL = PatternFill(start_color='B8CCE4', end_color='B8CCE4', fill_type='solid')
LIGHT_FILL = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')

# 对齐样式
CENTER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGN = Alignment(horizontal='left', vertical='center', wrap_text=True)
RIGHT_ALIGN = Alignment(horizontal='right', vertical='center')

# 数字格式
CURRENCY_FORMAT = '#,##0.00'
PERCENTAGE_FORMAT = '0.0000%'
NUMBER_FORMAT = '#,##0.0000'
DATE_FORMAT = 'YYYY-MM-DD'

# 预定义单元格样式组合
TITLE_STYLE = {
    'font': TITLE_FONT,
    'alignment': CENTER_ALIGN,
    'fill': TITLE_FILL,
    'border': THIN_BORDER,
}

HEADER_STYLE = {
    'font': HEADER_FONT,
    'alignment': CENTER_ALIGN,
    'fill': HEADER_FILL,
    'border': THIN_BORDER,
}

DATA_STYLE = {
    'font': DATA_FONT,
    'alignment': RIGHT_ALIGN,
    'border': THIN_BORDER,
}

DATA_STYLE_LEFT = {
    'font': DATA_FONT,
    'alignment': LEFT_ALIGN,
    'border': THIN_BORDER,
}

LABEL_STYLE = {
    'font': LABEL_FONT,
    'alignment': LEFT_ALIGN,
    'border': THIN_BORDER,
}


def apply_style(cell, style_dict):
    """将样式字典应用到单元格。"""
    for key, value in style_dict.items():
        setattr(cell, key, value)


def set_column_widths(ws, widths):
    """设置列宽。widths 是 dict: {col_letter: width}。"""
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
